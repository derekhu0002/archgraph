#!/usr/bin/env python3
"""Build the five-page ArchGraph coding deck PPTX.

Each slide is a 16:9 dark page: the rendered diagram (docs/ppt/assets/pN-*.png)
is the full-bleed base layer, and every text node of the source SVG is re-emitted
as a native (selectable, editable) PowerPoint text box positioned exactly where
the SVG drew it. The PNG keeps lines/badges/boxes faithful; the text boxes keep
the text editable and searchable.

    python scripts/build-coding-deck-pptx.py

Inputs : docs/ppt/content-plan.json, docs/ppt/assets/pN-*.svg, docs/ppt/assets/pN-*.png
Output : docs/ppt/archgraph-coding-business-deck.pptx
"""

from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "docs" / "ppt" / "assets"
OUT = ROOT / "docs" / "ppt" / "archgraph-coding-business-deck.pptx"
PLAN = json.loads((ROOT / "docs" / "ppt" / "content-plan.json").read_text(encoding="utf-8"))

SVG_NS = "{http://www.w3.org/2000/svg}"
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
EMU_PER_IN = 914400
SRC_W = 1280.0
SRC_H = 720.0
MAG = (SLIDE_W_IN * EMU_PER_IN) / SRC_W  # SVG user unit -> EMU


def hex_to_rgb(value: str) -> RGBColor:
    v = (value or "#eef1f6").lstrip("#")
    if len(v) == 3:
        v = "".join(c * 2 for c in v)
    return RGBColor(int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))


def collect_texts(svg_path: Path):
    """Return [{text,x,y,size,fill,anchor,weight,mono}] in source coordinates."""
    root = ET.parse(svg_path).getroot()
    out = []
    for node in root.iter(f"{SVG_NS}text"):
        raw = "".join(node.itertext())
        if not raw.strip():
            continue
        fam = node.get("font-family", "")
        out.append(
            {
                "text": raw,
                "x": float(node.get("x", 0)),
                "y": float(node.get("y", 0)),
                "size": float(node.get("font-size", 10)),
                "fill": node.get("fill", "#eef1f6"),
                "anchor": node.get("text-anchor", "start"),
                "weight": node.get("font-weight", "400"),
                "mono": "Consolas" in fam or "monospace" in fam,
            }
        )
    return out


def estimate_width(text: str, size: float) -> float:
    w = 0.0
    for ch in text:
        w += size if ord(ch) > 0x2E7F else size * 0.55
    return w


def add_text_box(slide, item):
    size = item["size"]
    text = item["text"]
    # generous box; PowerPoint never clips at the box edge for one line
    width = max(24.0, estimate_width(text, size) + 8.0)
    height = size * 1.7
    left = item["x"]
    if item["anchor"] == "middle":
        left -= width / 2.0
    elif item["anchor"] == "end":
        left -= width

    box = slide.shapes.add_textbox(
        Emu(int(left * MAG)),
        Emu(int((item["y"] - size * 1.16) * MAG)),
        Emu(int(width * MAG)),
        Emu(int(height * MAG)),
    )
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = text
    font = run.font
    font.size = Pt(size)
    font.bold = item["weight"] in ("600", "700", "800", "900", "bold")
    font.name = "Consolas" if item["mono"] else "Microsoft YaHei"
    font.color.rgb = hex_to_rgb(item["fill"])
    return box


def main() -> None:
    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W_IN * EMU_PER_IN))
    prs.slide_height = Emu(int(SLIDE_H_IN * EMU_PER_IN))
    blank = prs.slide_layouts[6]

    pages = sorted(ASSETS.glob("p[0-9]*-*.svg"))
    if not pages:
        raise SystemExit("no deck SVGs found; run scripts/gen-coding-deck.js first")

    for svg_path in pages:
        png_path = svg_path.with_suffix(".png")
        if not png_path.exists():
            raise SystemExit(f"missing rasterized PNG for {svg_path.name}; run scripts/rasterize-coding-deck.js")

        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(png_path), 0, 0, width=prs.slide_width, height=prs.slide_height
        )
        items = collect_texts(svg_path)
        for item in items:
            add_text_box(slide, item)
        print(f"slide {len(prs.slides._sldIdLst)}: {svg_path.name} + {len(items)} text boxes")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"wrote {OUT.relative_to(ROOT)} ({OUT.stat().st_size} bytes, {len(pages)} slides)")


if __name__ == "__main__":
    main()
